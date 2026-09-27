"""文件转换工具箱:图片 / PDF / Office 的本地转换(WPS 付费功能的免费本地替代)。

全部转换在本机完成;Office → PDF 优先 LibreOffice,其次调用本机
MS Office / WPS 的 COM 接口(自动探测);都不存在时给出安装指引。
"""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from .config import DATA_DIR

CONVERT_DIR = DATA_DIR / "convert"
MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_TOTAL_BYTES = 80 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_RENDER_PAGES = 100
RETENTION_SECONDS = 2 * 24 * 3600  # 临时文件保留 2 天

WORD_EXTS = {".doc", ".docx", ".rtf", ".txt", ".wps"}
EXCEL_EXTS = {".xls", ".xlsx", ".et", ".csv"}
PPT_EXTS = {".ppt", ".pptx", ".dps"}


class ToolError(Exception):
    """带用户可读中文提示的转换错误。"""


# ---------------------------------------------------------------- 基础工具
def _require_pillow():
    try:
        from PIL import Image
    except ImportError as exc:  # pragma: no cover
        raise ToolError("缺少 Pillow 依赖,请重新运行 scripts\\setup.bat 安装。") from exc
    return Image


def _require_pypdf():
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:  # pragma: no cover
        raise ToolError("缺少 pypdf 依赖,请重新运行 scripts\\setup.bat 安装。") from exc
    return PdfReader, PdfWriter


def _safe_name(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\s]+', "_", str(name or "output"))
    return name.strip("_")[:60] or "output"


def cleanup_old_files() -> int:
    """删除过期的临时转换文件。"""
    if not CONVERT_DIR.exists():
        return 0
    deadline = time.time() - RETENTION_SECONDS
    removed = 0
    for f in CONVERT_DIR.iterdir():
        try:
            if f.is_file() and f.stat().st_mtime < deadline:
                f.unlink()
                removed += 1
        except OSError:
            continue
    return removed


def _save_output(name: str, data: bytes) -> Path:
    CONVERT_DIR.mkdir(parents=True, exist_ok=True)
    path = CONVERT_DIR / f"{int(time.time() * 1000)}_{_safe_name(name)}"
    path.write_bytes(data)
    return path


def _check_input(name: str, data: bytes) -> None:
    if not data:
        raise ToolError(f"文件 {name} 是空的。")
    if len(data) > MAX_FILE_BYTES:
        raise ToolError(f"文件 {name} 太大(超过 30MB)。")


def parse_ranges(spec: str, total: int) -> list[int]:
    """把 "1-3,5" 解析为 0 基页序号列表(去重保序);空串表示全部页。"""
    spec = (spec or "").strip()
    if not spec:
        return list(range(total))
    pages: list[int] = []
    for part in re.split(r"[,、;；\s]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if m:
            start, end = int(m.group(1)), int(m.group(2))
            if start > end:
                start, end = end, start
            pages.extend(range(start, end + 1))
        elif part.isdigit():
            pages.append(int(part))
        else:
            raise ToolError(f"页码格式不对:{part}(示例:1-3,5)")
    result: list[int] = []
    for p in pages:
        if 1 <= p <= total and p not in result:
            result.append(p)
    if not result:
        raise ToolError("页码超出范围或为空。")
    return [p - 1 for p in result]


# ---------------------------------------------------------------- 图片类
def images_to_pdf(files: list[tuple[str, bytes]]) -> list[tuple[str, bytes]]:
    """多张图片合并成一个 A4 PDF。"""
    Image = _require_pillow()
    pages = []
    dpi = 150
    a4 = (int(8.27 * dpi), int(11.69 * dpi))
    margin = int(0.4 * dpi)
    for name, data in files:
        _check_input(name, data)
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Exception as exc:  # noqa: BLE001
            raise ToolError(f"{name} 不是有效的图片:{exc}") from exc
        if img.mode in ("RGBA", "LA", "P"):
            img = img.convert("RGBA")
            bg = Image.new("RGB", img.size, (255, 255, 255))
            bg.paste(img, mask=img.split()[-1])
            img = bg
        elif img.mode != "RGB":
            img = img.convert("RGB")
        page = Image.new("RGB", a4, (255, 255, 255))
        scale = min((a4[0] - 2 * margin) / img.width, (a4[1] - 2 * margin) / img.height)
        if scale < 1:
            img = img.resize((max(1, int(img.width * scale)), max(1, int(img.height * scale))), Image.LANCZOS)
        page.paste(img, ((a4[0] - img.width) // 2, (a4[1] - img.height) // 2))
        pages.append(page)
    if not pages:
        raise ToolError("没有可转换的图片。")
    buffer = io.BytesIO()
    pages[0].save(buffer, format="PDF", resolution=dpi, save_all=True, append_images=pages[1:])
    return [("images.pdf", buffer.getvalue())]


def images_convert(files: list[tuple[str, bytes]], target: str = "jpg") -> list[tuple[str, bytes]]:
    """批量转换图片格式(jpg / png / webp)。"""
    Image = _require_pillow()
    target = (target or "jpg").lower()
    if target not in ("jpg", "jpeg", "png", "webp"):
        raise ToolError(f"不支持的目标格式:{target}")
    fmt = "JPEG" if target in ("jpg", "jpeg") else target.upper()
    outputs = []
    for name, data in files:
        _check_input(name, data)
        try:
            img = Image.open(io.BytesIO(data))
            if getattr(img, "is_animated", False):
                img.seek(0)
            img.load()
        except Exception as exc:  # noqa: BLE001
            raise ToolError(f"{name} 不是有效的图片:{exc}") from exc
        if fmt == "JPEG":
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                bg = Image.new("RGB", img.size, (255, 255, 255))
                bg.paste(img, mask=img.split()[-1])
                img = bg
            elif img.mode != "RGB":
                img = img.convert("RGB")
        elif img.mode == "P":
            img = img.convert("RGBA")
        buffer = io.BytesIO()
        save_kwargs = {"quality": 90} if fmt == "JPEG" else {}
        img.save(buffer, format=fmt, **save_kwargs)
        stem = Path(name).stem
        ext = "jpg" if fmt == "JPEG" else fmt.lower()
        outputs.append((f"{stem}.{ext}", buffer.getvalue()))
    return outputs


def images_compress(files: list[tuple[str, bytes]], quality: int = 75, max_dim: int = 1600) -> list[tuple[str, bytes]]:
    """批量压缩图片(统一输出 JPG)。"""
    Image = _require_pillow()
    quality = max(30, min(95, int(quality)))
    max_dim = max(400, min(4000, int(max_dim)))
    outputs = []
    for name, data in files:
        _check_input(name, data)
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Exception as exc:  # noqa: BLE001
            raise ToolError(f"{name} 不是有效的图片:{exc}") from exc
        if img.mode != "RGB":
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                bg = Image.new("RGB", img.size, (255, 255, 255))
                bg.paste(img, mask=img.split()[-1])
                img = bg
            else:
                img = img.convert("RGB")
        longest = max(img.size)
        if longest > max_dim:
            scale = max_dim / longest
            img = img.resize((max(1, int(img.width * scale)), max(1, int(img.height * scale))), Image.LANCZOS)
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=quality, optimize=True)
        outputs.append((f"{Path(name).stem}_compressed.jpg", buffer.getvalue()))
    return outputs


def images_concat(files: list[tuple[str, bytes]], direction: str = "vertical", gap: int = 12) -> list[tuple[str, bytes]]:
    """多张图片拼接成一张长图(纵向 / 横向)。"""
    Image = _require_pillow()
    imgs = []
    for name, data in files:
        _check_input(name, data)
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Exception as exc:  # noqa: BLE001
            raise ToolError(f"{name} 不是有效的图片:{exc}") from exc
        if img.mode != "RGB":
            img = img.convert("RGB")
        imgs.append(img)
    if not imgs:
        raise ToolError("没有可拼接的图片。")
    if direction == "horizontal":
        height = max(i.height for i in imgs)
        width = sum(i.width for i in imgs) + gap * (len(imgs) - 1)
        canvas = Image.new("RGB", (width, height), (255, 255, 255))
        x = 0
        for img in imgs:
            canvas.paste(img, (x, (height - img.height) // 2))
            x += img.width + gap
    else:
        width = max(i.width for i in imgs)
        height = sum(i.height for i in imgs) + gap * (len(imgs) - 1)
        canvas = Image.new("RGB", (width, height), (255, 255, 255))
        y = 0
        for img in imgs:
            canvas.paste(img, ((width - img.width) // 2, y))
            y += img.height + gap
    buffer = io.BytesIO()
    canvas.save(buffer, format="PNG")
    return [("merged.png", buffer.getvalue())]


# ---------------------------------------------------------------- PDF 类
def _read_pdf(name: str, data: bytes):
    PdfReader, _ = _require_pypdf()
    _check_input(name, data)
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ToolError("这个 PDF 有密码,请先用「PDF 解密」转换。")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ToolError(f"PDF 页数太多(超过 {MAX_PDF_PAGES} 页)。")
        return reader
    except ToolError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ToolError(f"无法读取 PDF:{exc}") from exc


def _write_pdf(writer) -> bytes:
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def pdf_merge(files: list[tuple[str, bytes]]) -> list[tuple[str, bytes]]:
    """合并多个 PDF。"""
    _, PdfWriter = _require_pypdf()
    writer = PdfWriter()
    for name, data in files:
        reader = _read_pdf(name, data)
        for page in reader.pages:
            writer.add_page(page)
    return [("merged.pdf", _write_pdf(writer))]


def pdf_extract(files: list[tuple[str, bytes]], ranges: str = "") -> list[tuple[str, bytes]]:
    """提取指定页合并为新的 PDF。"""
    _, PdfWriter = _require_pypdf()
    name, data = files[0]
    reader = _read_pdf(name, data)
    pages = parse_ranges(ranges, len(reader.pages))
    writer = PdfWriter()
    for i in pages:
        writer.add_page(reader.pages[i])
    return [(f"{Path(name).stem}_extract.pdf", _write_pdf(writer))]


def pdf_split(files: list[tuple[str, bytes]], ranges: str = "") -> list[tuple[str, bytes]]:
    """拆分 PDF:按 ranges 每段一个文件;留空则每页一个文件。"""
    _, PdfWriter = _require_pypdf()
    name, data = files[0]
    reader = _read_pdf(name, data)
    total = len(reader.pages)
    stem = Path(name).stem
    outputs = []
    if (ranges or "").strip():
        groups = []
        for part in re.split(r"[,、;；\s]+", ranges.strip()):
            if part:
                groups.append(parse_ranges(part, total))
        if not groups:
            raise ToolError("拆分范围为空。")
        for idx, pages in enumerate(groups, 1):
            writer = PdfWriter()
            for i in pages:
                writer.add_page(reader.pages[i])
            outputs.append((f"{stem}_part{idx}.pdf", _write_pdf(writer)))
    else:
        if total > 50:
            raise ToolError("页数较多(超过 50 页),请填写拆分范围(如 1-10,11-20)。")
        for i in range(total):
            writer = PdfWriter()
            writer.add_page(reader.pages[i])
            outputs.append((f"{stem}_page{i + 1}.pdf", _write_pdf(writer)))
    return outputs


def pdf_rotate(files: list[tuple[str, bytes]], degrees: int = 90, ranges: str = "") -> list[tuple[str, bytes]]:
    """旋转指定页面(留空=全部)。"""
    _, PdfWriter = _require_pypdf()
    name, data = files[0]
    reader = _read_pdf(name, data)
    pages = parse_ranges(ranges, len(reader.pages))
    degrees = int(degrees) % 360
    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i in pages:
            page.rotate(degrees)
        writer.add_page(page)
    return [(f"{Path(name).stem}_rotated.pdf", _write_pdf(writer))]


def pdf_encrypt(files: list[tuple[str, bytes]], password: str = "") -> list[tuple[str, bytes]]:
    """给 PDF 设置密码。"""
    _, PdfWriter = _require_pypdf()
    name, data = files[0]
    password = (password or "").strip()
    if not password:
        raise ToolError("请填写要设置的密码。")
    reader = _read_pdf(name, data)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(password)
    return [(f"{Path(name).stem}_encrypted.pdf", _write_pdf(writer))]


def pdf_decrypt(files: list[tuple[str, bytes]], password: str = "") -> list[tuple[str, bytes]]:
    """移除 PDF 密码。"""
    PdfReader, PdfWriter = _require_pypdf()
    name, data = files[0]
    _check_input(name, data)
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            if not reader.decrypt(password or ""):
                raise ToolError("密码不正确,无法解密。")
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        return [(f"{Path(name).stem}_decrypted.pdf", _write_pdf(writer))]
    except ToolError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ToolError(f"解密失败:{exc}") from exc


def pdf_to_images(files: list[tuple[str, bytes]], fmt: str = "png", dpi: int = 150) -> list[tuple[str, bytes]]:
    """PDF 逐页导出为图片(最多前 100 页)。"""
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:  # pragma: no cover
        raise ToolError("缺少 pypdfium2 依赖,请重新运行 scripts\\setup.bat 安装。") from exc
    name, data = files[0]
    fmt = (fmt or "png").lower()
    if fmt not in ("png", "jpg", "jpeg"):
        raise ToolError(f"不支持的图片格式:{fmt}")
    if fmt in ("jpg", "jpeg"):
        fmt = "jpg"
    dpi = max(72, min(300, int(dpi)))
    _check_input(name, data)
    stem = Path(name).stem
    outputs = []
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.pdf"
        src.write_bytes(data)
        try:
            doc = pdfium.PdfDocument(str(src))
        except Exception as exc:  # noqa: BLE001
            raise ToolError(f"无法读取 PDF:{exc}") from exc
        try:
            count = min(len(doc), MAX_RENDER_PAGES)
            if count == 0:
                raise ToolError("这个 PDF 没有页面。")
            for i in range(count):
                page = doc[i]
                bitmap = page.render(scale=dpi / 72)
                pil = bitmap.to_pil()
                buffer = io.BytesIO()
                if fmt == "jpg":
                    pil.convert("RGB").save(buffer, format="JPEG", quality=88)
                else:
                    pil.save(buffer, format="PNG")
                outputs.append((f"{stem}_page{i + 1}.{fmt}", buffer.getvalue()))
        finally:
            doc.close()
    return outputs


def pdf_to_docx(files: list[tuple[str, bytes]]) -> list[tuple[str, bytes]]:
    """PDF 转 Word(纯文本版,不保留排版)。"""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ToolError("缺少 python-docx 依赖,请重新运行 scripts\\setup.bat 安装。") from exc
    name, data = files[0]
    reader = _read_pdf(name, data)
    document = docx.Document()
    has_text = False
    for idx, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        text = re.sub(r"[ \t]+", " ", text).strip()
        if text:
            has_text = True
            if idx > 1:
                document.add_page_break()
            for line in text.split("\n"):
                document.add_paragraph(line.strip())
    if not has_text:
        raise ToolError("这个 PDF 提取不到文字(可能是扫描版图片),无法转 Word。")
    buffer = io.BytesIO()
    document.save(buffer)
    return [(f"{Path(name).stem}.docx", buffer.getvalue())]


# ---------------------------------------------------------------- Office → PDF
def _ps_quote(text: str) -> str:
    return "'" + str(text).replace("'", "''") + "'"


def _find_soffice() -> str | None:
    candidates = []
    for name in ("soffice", "soffice.exe"):
        found = shutil.which(name)
        if found:
            candidates.append(found)
    candidates += [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for c in candidates:
        if c and Path(c).exists():
            return c
    return None


def _convert_with_libreoffice(src: Path, out_dir: Path) -> Path | None:
    soffice = _find_soffice()
    if not soffice:
        return None
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
            capture_output=True, timeout=180,
        )
    except (subprocess.TimeoutExpired, OSError):
        return None
    target = out_dir / (src.stem + ".pdf")
    return target if target.exists() else None


def _convert_with_com(src: Path, dst: Path) -> str:
    """用 PowerShell 调用 MS Office / WPS 的 COM 接口转换。返回 '' 表示成功,否则错误信息。"""
    ext = src.suffix.lower()
    if ext in EXCEL_EXTS:
        progids = "('Excel.Application', 'KET.Application')"
        script = f"""
$ErrorActionPreference = 'Stop'
$src = {_ps_quote(str(src))}
$dst = {_ps_quote(str(dst))}
$app = $null
foreach ($progid in {progids}) {{
  try {{ $app = New-Object -ComObject $progid; break }} catch {{ }}
}}
if ($null -eq $app) {{ Write-Output 'NOOFFICE'; exit 3 }}
try {{
  $app.Visible = $false
  $wb = $app.Workbooks.Open($src, 0, $true)
  try {{ $wb.ExportAsFixedFormat(0, $dst) }} finally {{ $wb.Close($false) }}
  Write-Output 'OK'
}} finally {{ $app.Quit() }}
"""
    elif ext in PPT_EXTS:
        progids = "('PowerPoint.Application', 'KWPP.Application')"
        script = f"""
$ErrorActionPreference = 'Stop'
$src = {_ps_quote(str(src))}
$dst = {_ps_quote(str(dst))}
$app = $null
foreach ($progid in {progids}) {{
  try {{ $app = New-Object -ComObject $progid; break }} catch {{ }}
}}
if ($null -eq $app) {{ Write-Output 'NOOFFICE'; exit 3 }}
try {{
  $pres = $app.Presentations.Open($src, $true, $false, $false)
  try {{ $pres.SaveAs($dst, 32) }} finally {{ $pres.Close() }}
  Write-Output 'OK'
}} finally {{ $app.Quit() }}
"""
    else:
        progids = "('Word.Application', 'KWPS.Application')"
        script = f"""
$ErrorActionPreference = 'Stop'
$src = {_ps_quote(str(src))}
$dst = {_ps_quote(str(dst))}
$app = $null
foreach ($progid in {progids}) {{
  try {{ $app = New-Object -ComObject $progid; break }} catch {{ }}
}}
if ($null -eq $app) {{ Write-Output 'NOOFFICE'; exit 3 }}
try {{
  $app.Visible = $false
  $doc = $app.Documents.Open($src, $false, $true)
  try {{ $doc.SaveAs([ref]$dst, [ref]17) }} finally {{ $doc.Close($false) }}
  Write-Output 'OK'
}} finally {{ $app.Quit() }}
"""
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
            capture_output=True, text=True, timeout=180,
        )
    except (subprocess.TimeoutExpired, OSError) as exc:
        return f"调用 Office 组件超时或失败:{exc}"
    stdout = (result.stdout or "").strip()
    if "NOOFFICE" in stdout:
        return "NOOFFICE"
    if result.returncode != 0:
        detail = (result.stderr or result.stdout or "").strip().splitlines()
        detail = detail[0] if detail else f"退出码 {result.returncode}"
        return f"Office 组件转换失败:{detail}"
    return ""


def office_to_pdf(files: list[tuple[str, bytes]]) -> list[tuple[str, bytes]]:
    """Word / Excel / PPT → PDF。优先 LibreOffice,其次本机 Office / WPS。"""
    name, data = files[0]
    _check_input(name, data)
    ext = Path(name).suffix.lower()
    if ext not in WORD_EXTS | EXCEL_EXTS | PPT_EXTS:
        raise ToolError(f"暂不支持这种格式:{ext or '(无扩展名)'}。支持 Word / Excel / PPT / txt。")
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        src = tmp_dir / f"input{ext}"
        src.write_bytes(data)

        pdf_path = _convert_with_libreoffice(src, tmp_dir)
        if pdf_path is None:
            dst = tmp_dir / "output.pdf"
            error = _convert_with_com(src, dst)
            if error == "NOOFFICE":
                raise ToolError(
                    "没有找到可用的转换组件:请安装 WPS、Microsoft Office 或 LibreOffice(免费)其中之一后重试。"
                )
            if error:
                raise ToolError(error)
            pdf_path = dst

        if pdf_path is None or not pdf_path.exists():
            raise ToolError("转换没有生成 PDF,请检查文件是否损坏。")
        return [(f"{Path(name).stem}.pdf", pdf_path.read_bytes())]


# ---------------------------------------------------------------- 总入口
def run_conversion(kind: str, files: list[tuple[str, bytes]], options: dict | None = None) -> list[tuple[str, Path]]:
    """执行转换,返回 [(文件名, 保存路径)]。"""
    options = options or {}
    cleanup_old_files()

    total = sum(len(data) for _, data in files)
    if total > MAX_TOTAL_BYTES:
        raise ToolError("上传的文件总大小超过 80MB,请分批转换。")
    if not files:
        raise ToolError("请先选择要转换的文件。")

    if kind == "images_to_pdf":
        outputs = images_to_pdf(files)
    elif kind == "image_convert":
        outputs = images_convert(files, str(options.get("target") or "jpg"))
    elif kind == "image_compress":
        outputs = images_compress(
            files,
            int(options.get("quality") or 75),
            int(options.get("max_dim") or 1600),
        )
    elif kind == "images_concat":
        outputs = images_concat(files, str(options.get("direction") or "vertical"))
    elif kind == "pdf_merge":
        outputs = pdf_merge(files)
    elif kind == "pdf_extract":
        outputs = pdf_extract(files, str(options.get("ranges") or ""))
    elif kind == "pdf_split":
        outputs = pdf_split(files, str(options.get("ranges") or ""))
    elif kind == "pdf_rotate":
        outputs = pdf_rotate(
            files,
            int(options.get("degrees") or 90),
            str(options.get("ranges") or ""),
        )
    elif kind == "pdf_encrypt":
        outputs = pdf_encrypt(files, str(options.get("password") or ""))
    elif kind == "pdf_decrypt":
        outputs = pdf_decrypt(files, str(options.get("password") or ""))
    elif kind == "pdf_to_images":
        outputs = pdf_to_images(
            files,
            str(options.get("fmt") or "png"),
            int(options.get("dpi") or 150),
        )
    elif kind == "pdf_to_docx":
        outputs = pdf_to_docx(files)
    elif kind == "office_to_pdf":
        outputs = office_to_pdf(files)
    else:
        raise ToolError(f"未知的转换类型:{kind}")

    saved: list[tuple[str, Path]] = []
    for name, data in outputs:
        saved.append((name, _save_output(name, data)))
    return saved
